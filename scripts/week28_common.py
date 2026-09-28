"""Week 28 IG shared paths and TFC reads.

Day 1 and Day 2 use the same window: start 2026-09-22, end 2026-09-25.
Three days of counts. Clocks stay as recorded.
"""
from __future__ import annotations

import re
import sys
from datetime import date, datetime
from pathlib import Path

APP = Path(__file__).resolve().parents[1]
SCRIPTS = Path(__file__).resolve().parent
sys.path.insert(0, str(APP))
sys.path.insert(0, str(SCRIPTS))

import openpyxl  # noqa: E402

FOLDER = Path(r"C:\Users\isaac\Downloads\week 28\week 28")
TFC_NAME = "Week 28 IG TFC.xlsx"
BACKUP_NAME = "_week28_backup"
PHOTO_BACKUP = "_week28_photo_backup"

DAY1_START = date(2026, 9, 22)
DAY1_END = date(2026, 9, 25)
DAY2_START = date(2026, 9, 22)
DAY2_END = date(2026, 9, 25)

DAY1_SHEET = "Week 28 issac day 1 "
DAY2_SHEET = "Week 28 issac day 2 "


def _cell_id(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].isdigit():
        return s[:-2]
    return s


def _serial_key(v) -> int | None:
    s = _cell_id(v)
    if not s or not s.isdigit():
        return None
    return int(s)


def serial_keys(v) -> set[int]:
    n = _serial_key(v)
    if n is None:
        return set()
    keys = {n}
    if n < 100000:
        keys.add(200000 + n)
    return keys


def _sheet(wb, want: str, *, day: int):
    if want in wb.sheetnames:
        return wb[want]
    needle = f"day {day}"
    for name in wb.sheetnames:
        if needle in name.lower():
            return wb[name]
    raise ValueError(f"missing Day {day} sheet; have {wb.sheetnames}")


def _installed(v) -> bool:
    return str(v or "").strip().lower() in {"x", "1", "true", "yes"}


def load_tfc(tfc: Path) -> dict:
    """Sites, installed flags, and 60min serial -> '{site}{dir}' for both days."""
    wb = openpyxl.load_workbook(tfc, data_only=True)
    ws1 = _sheet(wb, DAY1_SHEET, day=1)
    ws2 = _sheet(wb, DAY2_SHEET, day=2)
    day1: set[str] = set()
    day2: set[str] = set()
    day1_installed: set[str] = set()
    day2_installed: set[str] = set()
    serial_to_tag: dict[int, str] = {}
    order1: list[tuple[str, str]] = []
    order2: list[tuple[str, str]] = []
    gps1: dict[str, tuple[float, float]] = {}
    gps2: dict[str, tuple[float, float]] = {}
    missing_dir: list[str] = []

    def take(ws, day_sites: set[str], installed: set[str], order: list, gps: dict, day: int):
        for row in range(2, ws.max_row + 1):
            site = _cell_id(ws.cell(row, 2).value)
            if not site:
                continue
            day_sites.add(site)
            dirc = str(ws.cell(row, 4).value or "").strip().lower()[:1]
            inst = _installed(ws.cell(row, 7).value)
            skip = _installed(ws.cell(row, 8).value)
            lat, lon = ws.cell(row, 10).value, ws.cell(row, 11).value
            if isinstance(lat, (int, float)) and isinstance(lon, (int, float)):
                gps[site] = (float(lat), float(lon))
            if inst:
                installed.add(site)
            if inst and not skip:
                order.append((site, dirc))
            serial = ws.cell(row, 3).value
            if not inst or _serial_key(serial) is None:
                continue
            if dirc not in {"n", "e", "s", "w"}:
                missing_dir.append(f"day {day} site {site} serial {serial} dir={dirc!r}")
                continue
            tag = f"{site}{dirc}"
            for key in serial_keys(serial):
                if key in serial_to_tag and serial_to_tag[key] != tag:
                    raise ValueError(
                        f"serial {key} maps to both {serial_to_tag[key]} and {tag}"
                    )
                serial_to_tag[key] = tag

    take(ws1, day1, day1_installed, order1, gps1, 1)
    take(ws2, day2, day2_installed, order2, gps2, 2)
    if missing_dir:
        raise ValueError("rows missing direction: " + "; ".join(missing_dir))
    both_installed = day1_installed & day2_installed
    if both_installed:
        raise ValueError(f"site installed on both TFC days: {sorted(both_installed)}")
    return {
        "day1": day1,
        "day2": day2,
        "day1_installed": day1_installed,
        "day2_installed": day2_installed,
        "serial_to_tag": serial_to_tag,
        "order1": order1,
        "order2": order2,
        "gps1": gps1,
        "gps2": gps2,
        "sheet1": ws1.title,
        "sheet2": ws2.title,
    }


def resolve_count_site(name: str, known: set[str], header: bytes = b"") -> str:
    """Site id from a count filename, then from the TVP header name.

    50770nc1b.tvp is the file for TFC site 5070 (header text is 5070nc1b,
    serial 2208). A filename that already matches an installed site wins,
    so 4925nc1b.tvp stays 4925 even if the counter's location string differs.
    """
    stem = name.split(".")[0]
    m = re.match(r"(\d+)([nesw])c1b$", stem, re.I)
    raw = m.group(1) if m else ""
    if not raw:
        m2 = re.match(r"(\d+)", name)
        raw = m2.group(1) if m2 else ""
    if raw in known:
        return raw
    hm = re.search(rb"(\d{3,5})[nesw]c1b", header[:500], re.I)
    if hm and hm.group(1).decode() in known:
        return hm.group(1).decode()
    for sid in sorted(known, key=len, reverse=True):
        if raw.startswith(sid) and raw[len(sid) :] != "" and set(raw[len(sid) :]) <= {"0"}:
            return sid
    return raw


def as_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        base = datetime(1899, 12, 30)
        return (base + __import__("datetime").timedelta(days=float(v))).date()
    return None
