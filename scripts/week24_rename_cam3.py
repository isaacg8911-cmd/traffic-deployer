"""Prefix Week 24 Day 1 Cam 3 CSVs with TFC site + direction.

CSV names are Traficam serials (e.g. 001712-118.60min.csv). TFC Day 1 rows
with Notes = cam 3 map serial -> site + n/e. Writes 3997e-001712-118.60min.csv.

Default folder: C:\\Users\\isaac\\Downloads\\WEEK 24 1
Does not modify the source week-24-ig folder.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl

TFC = Path(r"C:\Users\isaac\Downloads\week 24 ig (2)\week 24 ig\Week 24 IG TFC.xlsx")
OUT = Path(r"C:\Users\isaac\Downloads\WEEK 24 1")
DAY1_SHEET = "Week 24 Day 1 Isaac"


def _cell_id(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].isdigit():
        return s[:-2]
    return s


def load_cam3_serial_to_site(tfc: Path) -> dict[str, str]:
    """serial -> '{site}{n|e}' tag from Day 1 cam 3 TFC rows."""
    wb = openpyxl.load_workbook(tfc, data_only=True)
    if DAY1_SHEET not in wb.sheetnames:
        raise ValueError(f"missing sheet {DAY1_SHEET!r}; have {wb.sheetnames}")
    ws = wb[DAY1_SHEET]
    mapping: dict[str, str] = {}
    missing_dir: list[str] = []
    for row in range(2, ws.max_row + 1):
        notes = str(ws.cell(row, 6).value or "").strip().lower()
        if "cam 3" not in notes:
            continue
        site = _cell_id(ws.cell(row, 2).value)
        serial = _cell_id(ws.cell(row, 3).value)
        dirc = str(ws.cell(row, 4).value or "").strip().lower()[:1]
        if not site or not serial:
            continue
        if dirc not in {"n", "e"}:
            missing_dir.append(f"site {site} serial {serial} dir={dirc!r}")
            continue
        tag = f"{site}{dirc}"
        if serial in mapping and mapping[serial] != tag:
            raise ValueError(f"serial {serial} maps to both {mapping[serial]} and {tag}")
        mapping[serial] = tag
    if missing_dir:
        raise ValueError("cam 3 rows missing n/e direction: " + "; ".join(missing_dir))
    return mapping


def file_serial(name: str) -> str | None:
    """Serial from 001712-…, 3997-001712-…, or 3997e-001712-…."""
    stem = name[:-4] if name.lower().endswith(".csv") else name
    if not stem:
        return None
    first = stem.split("-", 1)[0]
    if first.isdigit() and len(first) >= 6:
        return first
    m = re.match(r"^(\d{3,5})[nsew]?-(.+)$", stem, re.I)
    if not m:
        return None
    rest_first = m.group(2).split("-", 1)[0]
    if rest_first.isdigit() and len(rest_first) >= 6:
        return rest_first
    return None


def dest_name(name: str, serial_to_site: dict[str, str]) -> str | None:
    serial = file_serial(name)
    if not serial:
        return None
    tag = serial_to_site.get(serial)
    if not tag:
        return None
    prefix = f"{tag}-"
    if name.startswith(prefix):
        return name
    m = re.match(r"^\d{3,5}[nsew]?-(.+)$", name, re.I)
    if m and m.group(1).startswith(serial):
        return prefix + m.group(1)
    return prefix + name


def rename_dir(folder: Path, serial_to_site: dict[str, str]) -> tuple[list[tuple[str, str]], list[str], list[str]]:
    renamed: list[tuple[str, str]] = []
    skipped: list[str] = []
    errors: list[str] = []
    csvs = sorted(p for p in folder.iterdir() if p.is_file() and p.suffix.lower() == ".csv")
    for src in csvs:
        new_name = dest_name(src.name, serial_to_site)
        if new_name is None:
            continue
        if new_name == src.name:
            skipped.append(src.name)
            continue
        dest = src.with_name(new_name)
        if dest.exists():
            errors.append(f"dest exists: {src.name} -> {new_name}")
            continue
        src.rename(dest)
        renamed.append((src.name, new_name))
    return renamed, skipped, errors


def main() -> int:
    folder = Path(sys.argv[1]) if len(sys.argv) > 1 else OUT
    tfc = Path(sys.argv[2]) if len(sys.argv) > 2 else TFC
    if not folder.is_dir():
        print(f"missing folder {folder}", file=sys.stderr)
        return 1
    if not tfc.is_file():
        print(f"missing TFC {tfc}", file=sys.stderr)
        return 1

    mapping = load_cam3_serial_to_site(tfc)
    print(f"TFC cam 3 serials: {len(mapping)}")
    renamed, skipped, errors = rename_dir(folder, mapping)

    used = {file_serial(b) or file_serial(a) for a, b in renamed}
    used |= {file_serial(n) for n in skipped}
    unmatched_tfc = sorted(
        (ser, site) for ser, site in mapping.items() if ser not in used
    )
    unmatched_files = []
    for p in sorted(folder.glob("*.60min.csv")):
        ser = file_serial(p.name)
        if ser and ser not in mapping:
            unmatched_files.append(p.name)

    for a, b in renamed:
        print(f"  {a} -> {b}")
    if skipped:
        print(f"already prefixed: {len(skipped)}")
        for n in skipped:
            print(f"  {n}")
    print(f"renamed {len(renamed)}, already {len(skipped)}, errors {len(errors)}")
    if unmatched_tfc:
        print("TFC cam 3 with no CSV (expected if unit failed/off):")
        for ser, site in unmatched_tfc:
            print(f"  site {site} serial {ser}")
    if unmatched_files:
        print("60min CSV serial not on TFC cam 3:")
        for n in unmatched_files:
            print(f"  {n}")
    for e in errors:
        print(f"FAIL {e}", file=sys.stderr)
    if errors or unmatched_files:
        return 1
    if not renamed and not skipped:
        print("no Day 1 cam 3 CSVs found", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
