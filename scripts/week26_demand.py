"""Week 26 IG: restamp 3 days of real data onto Director timeline.

Day 2 (TFC Day 2 sites): start 2026-09-07, end 2026-09-10 — volume CSVs,
60min CSVs, hose TVPs. Keep clock times. Drop rows/hits past the window.

Day 1 (TFC Day 1 sites): start 2026-09-08, TVP end Friday 2026-09-11 —
same rule. TFC Date column set to that day's start.

60min CSVs: match Traficam serial to TFC; prefix '{site}{n|e} '.

Default folder: C:\\Users\\isaac\\Downloads\\week 26 ig\\week 26 ig
Backs up data files to _week26_backup/ then writes in place.
"""
from __future__ import annotations

import csv
import re
import shutil
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

APP = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import openpyxl  # noqa: E402
from core import picocount_hits  # noqa: E402
from week24_trim import (  # noqa: E402
    _join_like,
    locate_stream,
    parse_60min,
    parse_volume_csv,
    rewrite_tvp,
    verify_tvp,
    write_volume_csv,
)
from week25_day2_rename_day1_shift import (  # noqa: E402
    dest_name,
    file_serial_key,
)

FOLDER = Path(r"C:\Users\isaac\Downloads\week 26 ig\week 26 ig")
TFC_NAME = "Week 26 IG TFC.xlsx"
BACKUP_NAME = "_week26_backup"

DAY1_START = date(2026, 9, 8)
DAY1_END = date(2026, 9, 11)  # Friday
DAY2_START = date(2026, 9, 7)
DAY2_END = date(2026, 9, 10)

DAY1_SHEET = "Week 26 Day 1 issac"
DAY2_SHEET = "Week 26 Day 2  issac"


def _cell_id(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].isdigit():
        return s[:-2]
    return s


def _serial_key(v) -> int | None:
    s = _cell_id(v)
    if not s or not s.isdigit():
        return None
    return int(s)


def _sheet(wb, want: str, *, day: int):
    if want in wb.sheetnames:
        return wb[want]
    needle = f"day {day}"
    for name in wb.sheetnames:
        if needle in name.lower():
            return wb[name]
    raise ValueError(f"missing Day {day} sheet; have {wb.sheetnames}")


def load_tfc_sites(tfc: Path) -> tuple[set[str], set[str], dict[int, str]]:
    """Day 1 sites, Day 2 sites, serial_int -> '{site}{dir}' (Day 2 cam/60min)."""
    wb = openpyxl.load_workbook(tfc, data_only=True)
    ws1 = _sheet(wb, DAY1_SHEET, day=1)
    ws2 = _sheet(wb, DAY2_SHEET, day=2)
    day1: set[str] = set()
    day2: set[str] = set()
    serial_to_tag: dict[int, str] = {}
    missing_dir: list[str] = []
    for row in range(2, ws1.max_row + 1):
        site = _cell_id(ws1.cell(row, 2).value)
        if site:
            day1.add(site)
    for row in range(2, ws2.max_row + 1):
        site = _cell_id(ws2.cell(row, 2).value)
        serial = _serial_key(ws2.cell(row, 3).value)
        dirc = str(ws2.cell(row, 4).value or "").strip().lower()[:1]
        if site:
            day2.add(site)
        if not site or serial is None:
            continue
        if dirc not in {"n", "e", "s", "w"}:
            missing_dir.append(f"site {site} serial {serial} dir={dirc!r}")
            continue
        tag = f"{site}{dirc}"
        if serial in serial_to_tag and serial_to_tag[serial] != tag:
            raise ValueError(
                f"serial {serial} maps to both {serial_to_tag[serial]} and {tag}"
            )
        serial_to_tag[serial] = tag
    if missing_dir:
        raise ValueError("Day 2 rows missing direction: " + "; ".join(missing_dir))
    overlap = day1 & day2
    if overlap:
        raise ValueError(f"site on both TFC days: {sorted(overlap)}")
    return day1, day2, serial_to_tag


def site_of(name: str) -> str:
    m = re.match(r"(\d+)", name)
    return m.group(1) if m else ""


def stamp_tfc_dates(tfc: Path) -> dict[str, int]:
    wb = openpyxl.load_workbook(tfc)
    ws1 = _sheet(wb, DAY1_SHEET, day=1)
    ws2 = _sheet(wb, DAY2_SHEET, day=2)
    n1 = n2 = 0
    for row in range(2, ws1.max_row + 1):
        if ws1.cell(row, 1).value not in (None, ""):
            ws1.cell(row, 1).value = datetime(DAY1_START.year, DAY1_START.month, DAY1_START.day)
            n1 += 1
    for row in range(2, ws2.max_row + 1):
        if ws2.cell(row, 1).value not in (None, ""):
            ws2.cell(row, 1).value = datetime(DAY2_START.year, DAY2_START.month, DAY2_START.day)
            n2 += 1
    wb.save(tfc)
    return {"day1": n1, "day2": n2}


def _rewrite_60min_row(line: str, new_ts: datetime) -> str:
    parts = next(csv.reader([line]))
    if len(parts) < 2:
        raise ValueError(f"short 60min row: {line!r}")
    parts[0] = new_ts.strftime("%Y-%m-%d")
    parts[1] = new_ts.strftime("%H:%M")
    return _join_like(line, ["" if p is None else str(p) for p in parts])


def write_60min_shifted(
    path: Path,
    parsed: dict,
    *,
    new_start: datetime,
    cutoff: datetime,
    shift_days: int,
) -> dict:
    shift = timedelta(days=shift_days)
    kept: list[tuple[datetime, str]] = []
    for ts, line in parsed["rows"]:
        new_ts = ts - shift
        if new_ts >= cutoff:
            continue
        kept.append((new_ts, _rewrite_60min_row(line, new_ts)))
    dropped = len(parsed["rows"]) - len(kept)
    if not kept:
        raise ValueError(f"{path.name}: no 60min rows inside {new_start.date()}..{cutoff}")

    # Director window: end date + original end clock, even if the last
    # hourly bin is earlier (unit stopped; still stamp the study end).
    new_end = datetime.combine(cutoff.date(), parsed["end"].time())
    if new_end > cutoff:
        new_end = cutoff
    if new_end <= new_start:
        new_end = kept[-1][0] + timedelta(hours=1)

    duration_h = int(round((new_end - new_start).total_seconds() / 3600.0))
    if duration_h < 0:
        duration_h = 0
    total = 0
    for _ts, line in kept:
        parts = next(csv.reader([line]))
        for cell in parts[2:]:
            cell = cell.strip()
            if cell:
                try:
                    total += int(float(cell))
                except ValueError:
                    pass

    lines = list(parsed["lines"])
    out_lines: list[str] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("Study Start:"):
            orig_parts = next(csv.reader([line]))
            out_lines.append(
                _join_like(
                    line,
                    [orig_parts[0], new_start.strftime("%Y-%m-%d %H:%M:%S"), *orig_parts[2:]],
                )
            )
        elif line.startswith("Study End:"):
            orig_parts = next(csv.reader([line]))
            out_lines.append(
                _join_like(
                    line,
                    [orig_parts[0], new_end.strftime("%Y-%m-%d %H:%M:%S"), *orig_parts[2:]],
                )
            )
        elif line.startswith("Duration:"):
            orig_parts = next(csv.reader([line]))
            out_lines.append(_join_like(line, [orig_parts[0], str(duration_h), *orig_parts[2:]]))
        elif line.startswith("Total Vehicles:"):
            orig_parts = next(csv.reader([line]))
            out_lines.append(_join_like(line, [orig_parts[0], f" {total}", *orig_parts[2:]]))
        elif line.startswith("DATE,"):
            out_lines.append(line)
            for _ts, row in kept:
                out_lines.append(row)
            i += 1
            while i < len(lines) and re.match(r"^\d{4}-\d{2}-\d{2},", lines[i]):
                i += 1
            continue
        else:
            out_lines.append(line)
        i += 1

    nl = "\r\n" if parsed["newline"] == b"\r\n" else "\n"
    body = nl.join(out_lines) + nl
    data = body.encode("utf-8")
    if parsed["bom"]:
        data = b"\xef\xbb\xbf" + data
    path.write_bytes(data)
    return {
        "rows_in": len(parsed["rows"]),
        "rows_out": len(kept),
        "dropped": dropped,
        "new_start": new_start,
        "new_end": new_end,
        "orig_start": parsed["start"],
        "orig_end": parsed["end"],
        "total": total,
        "first": kept[0][0],
        "last": kept[-1][0],
        "shift_days": shift_days,
    }


def process_60min(src: Path, dst: Path, lines: list[str]) -> None:
    parsed = parse_60min(src)
    if not parsed["rows"]:
        raise ValueError(f"{src.name}: no data rows")
    first = parsed["rows"][0][0]
    shift_days = (first.date() - DAY2_START).days
    new_start = datetime.combine(DAY2_START, parsed["start"].time())
    cutoff = datetime.combine(DAY2_END, parsed["end"].time())
    info = write_60min_shifted(
        dst, parsed, new_start=new_start, cutoff=cutoff, shift_days=shift_days
    )
    check = parse_60min(dst)
    if check["start"].date() != DAY2_START:
        raise ValueError(f"{src.name}: start {check['start']} not on {DAY2_START}")
    if check["end"].date() > DAY2_END:
        raise ValueError(f"{src.name}: end {check['end']} after {DAY2_END}")
    bad = [ts for ts, _ in check["rows"] if ts >= cutoff or ts.date() > DAY2_END]
    if bad:
        raise ValueError(f"{src.name}: leftover rows after cutoff: {bad[:3]}")
    span_days = (check["rows"][-1][0].date() - check["rows"][0][0].date()).days + 1
    if span_days > 4:
        raise ValueError(f"{src.name}: kept {span_days} calendar days, want <= 4 (3-day count)")
    lines.append(
        f"CSV60 {src.name}: {parsed['start']}..{parsed['end']} -> "
        f"{info['new_start']}..{info['new_end']}; "
        f"rows {info['rows_in']} -> {info['rows_out']} (drop {info['dropped']}); "
        f"span {info['first']} .. {info['last']}"
    )


def process_volume(src: Path, dst: Path, lines: list[str]) -> None:
    parsed = parse_volume_csv(src)
    shift = (parsed["started"].date() - DAY2_START).days
    new_started = datetime.combine(DAY2_START, parsed["started"].time())
    new_ended = datetime.combine(DAY2_END, parsed["ended"].time())
    info = write_volume_csv(
        dst, parsed, new_started=new_started, new_ended=new_ended, shift_days=shift
    )
    check = parse_volume_csv(dst)
    if check["started"].date() != DAY2_START:
        raise ValueError(f"{src.name}: started {check['started']} not on {DAY2_START}")
    if check["ended"].date() > DAY2_END or check["ended"] > new_ended:
        raise ValueError(f"{src.name}: ended {check['ended']} after {new_ended}")
    for day in check["days"]:
        if day["date"] > DAY2_END:
            raise ValueError(f"{src.name}: day {day['date']} after {DAY2_END}")
        if day["date"] < DAY2_START:
            raise ValueError(f"{src.name}: day {day['date']} before {DAY2_START}")
    span = (check["days"][-1]["date"] - check["days"][0]["date"]).days + 1
    if span > 4:
        raise ValueError(f"{src.name}: kept {span} calendar days")
    lines.append(
        f"VOL {src.name}: {parsed['started']}..{parsed['ended']} -> "
        f"{info['new_started']}..{info['actual_end']}; "
        f"intervals {info['intervals_out']} drop {info['dropped_intervals']}; "
        f"grand {info['grand_total']}"
    )


def process_tvp(src: Path, dst: Path, *, start_day: date, end_day: date, lines: list[str]) -> None:
    data = src.read_bytes()
    orig_start = picocount_hits.decode_start_time_tvp(data)
    if orig_start is None:
        raise ValueError(f"{src.name}: no study start")
    shift = (orig_start.date() - start_day).days
    new_start = datetime.combine(start_day, orig_start.time())
    _off, hits, _end = locate_stream(data)
    valid = [h for h in hits if 0 <= h["seconds"] <= 86400 * 8]
    if not valid:
        raise ValueError(f"{src.name}: no valid hits")
    orig_last = orig_start + timedelta(seconds=valid[-1]["seconds"])
    shifted_last = orig_last - timedelta(days=shift)
    end_cap = datetime.combine(end_day, orig_last.time())
    cutoff = min(shifted_last, end_cap)
    info = rewrite_tvp(src, dst, new_start=new_start, cutoff=cutoff, orig_start=orig_start)
    v = verify_tvp(dst, expect_start=new_start, cutoff=cutoff)
    if v["start"].date() != start_day:
        raise ValueError(f"{src.name}: start {v['start']} not on {start_day}")
    if v["last"].date() > end_day:
        raise ValueError(f"{src.name}: last {v['last']} after {end_day}")
    lines.append(
        f"TVP {src.name}: start {orig_start} -> {v['start']}; "
        f"last {orig_last} -> {v['last']}; hits {info['hits_in']} -> {info['hits_out']} "
        f"(drop {info['dropped']}); bytes {info['bytes_in']} -> {info['bytes_out']}"
    )


def backup_file(src: Path, backup: Path) -> None:
    dest = backup / src.name
    if not dest.exists():
        shutil.copy2(src, dest)


def rename_60min(
    folder: Path, serial_to_tag: dict[int, str]
) -> tuple[list[tuple[str, str]], list[str], list[str]]:
    renamed: list[tuple[str, str]] = []
    skipped: list[str] = []
    errors: list[str] = []
    for src in sorted(folder.glob("*.60min.csv")):
        new_name = dest_name(src.name, serial_to_tag)
        if new_name is None:
            errors.append(f"no TFC match: {src.name}")
            continue
        if new_name == src.name:
            skipped.append(src.name)
            continue
        dest = src.with_name(new_name)
        if dest.exists():
            errors.append(f"dest exists: {src.name} -> {new_name}")
            continue
        src.rename(dest)
        renamed.append((src.name, new_name))
    return renamed, skipped, errors


def main() -> int:
    folder = Path(sys.argv[1]) if len(sys.argv) > 1 else FOLDER
    tfc = Path(sys.argv[2]) if len(sys.argv) > 2 else folder / TFC_NAME
    report_path = folder / "WEEK26_REPORT.txt"
    lines: list[str] = []

    def log(msg: str = "") -> None:
        print(msg)
        lines.append(msg)

    if not folder.is_dir():
        print(f"missing folder {folder}", file=sys.stderr)
        return 1
    if not tfc.is_file():
        print(f"missing TFC {tfc}", file=sys.stderr)
        return 1

    day1_sites, day2_sites, serial_to_tag = load_tfc_sites(tfc)
    backup = folder / BACKUP_NAME
    backup.mkdir(exist_ok=True)

    log(f"Folder: {folder}")
    log(f"TFC: {tfc}")
    log(f"Backup: {backup}")
    log(f"Day 2 window: {DAY2_START} .. {DAY2_END} (keep times, 3-day count)")
    log(f"Day 1 window: {DAY1_START} .. {DAY1_END} Friday (keep times, 3-day count)")
    log(f"TFC Day 1 sites: {len(day1_sites)}  Day 2 sites: {len(day2_sites)}  "
        f"Day 2 serials: {len(serial_to_tag)}")
    log()

    errors: list[str] = []
    n_ok = 0

    # --- TFC dates ---
    backup_file(tfc, backup)
    tfc_counts = stamp_tfc_dates(tfc)
    log(f"=== TFC dates ===")
    log(f"  Day 1 Date cells -> {DAY1_START}: {tfc_counts['day1']}")
    log(f"  Day 2 Date cells -> {DAY2_START}: {tfc_counts['day2']}")
    log()

    # --- 60min (Day 2 cam3) ---
    log("=== Day 2 60min CSVs ===")
    csv60 = sorted(folder.glob("*.60min.csv"))
    if not csv60:
        errors.append("no *.60min.csv files")
        log("  FAIL no *.60min.csv files")
    for src in csv60:
        backup_file(src, backup)
        try:
            n0 = len(lines)
            process_60min(src, src, lines)
            for msg in lines[n0:]:
                print(msg)
            n_ok += 1
        except Exception as exc:  # noqa: BLE001
            errors.append(f"{src.name}: {exc}")
            log(f"  FAIL {src.name}: {exc}")
            bak = backup / src.name
            if bak.exists():
                shutil.copy2(bak, src)
    log()

    # --- Day 2 volume CSVs ---
    log("=== Day 2 volume CSVs ===")
    vol_files = [
        p
        for p in sorted(folder.glob("*.csv"))
        if ".60min" not in p.name.lower() and site_of(p.name) in day2_sites
    ]
    if not vol_files:
        errors.append("no Day 2 volume CSVs")
        log("  FAIL no Day 2 volume CSVs")
    for src in vol_files:
        backup_file(src, backup)
        try:
            n0 = len(lines)
            process_volume(src, src, lines)
            for msg in lines[n0:]:
                print(msg)
            n_ok += 1
        except Exception as exc:  # noqa: BLE001
            errors.append(f"{src.name}: {exc}")
            log(f"  FAIL {src.name}: {exc}")
            bak = backup / src.name
            if bak.exists():
                shutil.copy2(bak, src)
    log()

    # --- TVPs ---
    log("=== TVP date shift + trim ===")
    tvps = sorted(folder.glob("*.tvp"))
    if not tvps:
        errors.append("no TVP files")
        log("  FAIL no TVP files")
    unknown_tvp: list[str] = []
    for src in tvps:
        site = site_of(src.name)
        if site in day1_sites:
            start_day, end_day = DAY1_START, DAY1_END
        elif site in day2_sites:
            start_day, end_day = DAY2_START, DAY2_END
        else:
            unknown_tvp.append(src.name)
            errors.append(f"TVP site not on TFC: {src.name}")
            log(f"  FAIL TVP site not on TFC: {src.name}")
            continue
        backup_file(src, backup)
        try:
            n0 = len(lines)
            process_tvp(src, src, start_day=start_day, end_day=end_day, lines=lines)
            for msg in lines[n0:]:
                print(msg)
            n_ok += 1
        except Exception as exc:  # noqa: BLE001
            errors.append(f"{src.name}: {exc}")
            log(f"  FAIL {src.name}: {exc}")
            bak = backup / src.name
            if bak.exists():
                shutil.copy2(bak, src)
    log()

    # --- rename 60min after content write ---
    log("=== Day 2 60min rename (serial -> site+dir) ===")
    renamed, skipped, rename_err = rename_60min(folder, serial_to_tag)
    for a, b in renamed:
        log(f"  {a} -> {b}")
    for n in skipped:
        log(f"  already: {n}")
    for e in rename_err:
        errors.append(e)
        log(f"  FAIL {e}")
    used_keys = {file_serial_key(b) or file_serial_key(a) for a, b in renamed}
    used_keys |= {file_serial_key(n) for n in skipped}
    wb = openpyxl.load_workbook(tfc, data_only=True)
    ws2 = _sheet(wb, DAY2_SHEET, day=2)
    cam3_missing: list[str] = []
    for row in range(2, ws2.max_row + 1):
        notes = str(ws2.cell(row, 6).value or "").lower()
        serial = _serial_key(ws2.cell(row, 3).value)
        site = _cell_id(ws2.cell(row, 2).value)
        if serial is None or "cam" not in notes:
            continue
        if serial not in used_keys:
            cam3_missing.append(f"  site {site} serial {serial}")
    if cam3_missing:
        log("TFC cam3 serials with no 60min CSV:")
        for c in cam3_missing:
            log(c)
    log(f"renamed={len(renamed)} already={len(skipped)} errors={len(rename_err)}")
    log()

    log("=== Summary ===")
    log(f"processed ok: {n_ok}")
    log(f"errors: {len(errors)}")
    if errors:
        for e in errors:
            log(f"  ERROR {e}")
    report_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    log(f"Report: {report_path}")
    if errors:
        return 1
    if not renamed and not skipped:
        print("no 60min CSVs renamed", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
