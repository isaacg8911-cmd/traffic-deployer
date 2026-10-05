"""Week 29 IG: Day 1 from Isaacx + Apple Maps, then 3-day clocks.

Day 1 start 2026-09-28, end 2026-10-01. Day 2 start 2026-09-30, end 2026-10-03.
Each file keeps its own clock and is cropped at that end time.

Day 1 TFC rows are the Apple guide, matched to Week 29 Isaacx.xls by street.
GPS is the Apple pin. Hose files supply site, direction, and counter serial.
"""
from __future__ import annotations

import base64
import json
import re
import shutil
import struct
import sys
from copy import copy
from datetime import date, datetime
from pathlib import Path

APP = Path(__file__).resolve().parents[1]
SCRIPTS = Path(__file__).resolve().parent
sys.path.insert(0, str(APP))
sys.path.insert(0, str(SCRIPTS))

import openpyxl  # noqa: E402
from core import ingest  # noqa: E402
from core.direction import infer_from_segment  # noqa: E402
from core.est_viewer import pushpins_from_est  # noqa: E402
from week17_segment_match import point_to_segment_m  # noqa: E402
from week24_trim import parse_60min  # noqa: E402
from week25_day2_rename_day1_shift import dest_name  # noqa: E402
from week26_photo_dates import (  # noqa: E402
    list_photos,
    read_datetime_original,
    set_file_times,
    shift_jpeg_dates,
)
from week27_est_gps import patch_map  # noqa: E402
from week28_step3_photos import (  # noqa: E402
    HIGH,
    LOW,
    _assign,
    compress_about_800kb,
)
import week26_demand as w26  # noqa: E402
from PIL import Image, ImageOps  # noqa: E402

FOLDER = Path(r"C:\Users\isaac\Downloads\week 29\week 29")
XLS = Path(r"C:\Users\isaac\Downloads\Week 29 Isaacx.xls")
EST_DAY1 = Path(r"C:\Users\isaac\Downloads\Week 29 Isaac day 1.est")
TFC_NAME = "Week 29 IG TFC.xlsx"
BACKUP_NAME = "_week29_backup"
PHOTO_BACKUP = "_week29_photo_backup"
GUIDE_JS = Path(r"C:\Users\isaac\AppData\Local\Temp\week29_script.js")

DAY1_START = date(2026, 9, 28)
DAY1_END = date(2026, 10, 1)
DAY2_START = date(2026, 9, 30)
DAY2_END = date(2026, 10, 3)

w26.DAY1_START = DAY1_START
w26.DAY1_END = DAY1_END
w26.DAY2_START = DAY2_START
w26.DAY2_END = DAY2_END

STREET_SUFFIX = re.compile(
    r"\b(STREET|ST|AVENUE|AVE|BOULEVARD|BLVD|DRIVE|DR|ROAD|RD|LANE|LN|COURT|CT|WAY)\b"
)


def street_key(text: str) -> str:
    s = re.sub(r"[^A-Za-z ]", " ", text).upper()
    s = STREET_SUFFIX.sub(" ", s)
    return " ".join(s.split())


def place_label(raw: str) -> str:
    m = re.search(
        r"([A-Za-z]+(?:\s+[A-Za-z]+)*)\s+"
        r"(St|Ave|Blvd|Dr|Rd|Ln|Ct|Way)\b",
        raw,
    )
    return m.group(0) if m else raw.strip()


def apple_places() -> list[dict]:
    script = json.loads(GUIDE_JS.read_text(encoding="utf-8"))
    ref = script["initialState"]["cards"][0]["ref"]
    blob = base64.b64decode(ref + "=" * ((4 - len(ref) % 4) % 4))
    strings = [(m.start(), m.group().decode()) for m in re.finditer(rb"[\x20-\x7e]{6,}", blob)]
    places = []
    for i in range(len(blob) - 8):
        lat = struct.unpack_from("<d", blob, i)[0]
        if not (33.0 < lat < 35.0):
            continue
        lon = None
        for j in range(i + 8, min(i + 32, len(blob) - 7)):
            cand = struct.unpack_from("<d", blob, j)[0]
            if -118.5 < cand < -116.0:
                lon = cand
                break
        if lon is None:
            continue
        raw = ""
        for soff, text in strings:
            if soff >= i:
                break
            if "Lake Elsinore" in text or "Leach" in text:
                raw = text
        places.append({"lat": lat, "lon": lon, "name": place_label(raw), "raw": raw})
    return places


def _cell_id(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].isdigit():
        return s[:-2]
    return s


def _installed(v) -> bool:
    return str(v or "").strip().lower() in {"x", "1", "true", "yes"}


def serial_keys(v) -> set[int]:
    s = _cell_id(v)
    if not s or not s.isdigit():
        return set()
    n = int(s)
    keys = {n}
    if n < 100000:
        keys.add(200000 + n)
    return keys


def tvp_header(path: Path) -> tuple[str, str]:
    """Return (8-digit serial, header location such as 4262nc1b)."""
    data = path.read_bytes()[:500]
    m = re.search(rb"(\d{8})\x00+(\d{4}[nesw]c1b)", data, re.I)
    if not m:
        raise ValueError(f"{path.name}: no serial/location in header")
    return m.group(1).decode(), m.group(2).decode().lower()


def load_day2(tfc: Path) -> dict:
    wb = openpyxl.load_workbook(tfc, data_only=True)
    ws = wb["Day 2"]
    installed: set[str] = set()
    serial_to_tag: dict[int, str] = {}
    order: list[tuple[str, str]] = []
    gps: dict[str, tuple[float, float]] = {}
    for row in range(2, ws.max_row + 1):
        site = _cell_id(ws.cell(row, 2).value)
        if not site:
            continue
        direc = str(ws.cell(row, 4).value or "").strip().lower()[:1]
        if _installed(ws.cell(row, 7).value):
            installed.add(site)
            if not _installed(ws.cell(row, 8).value):
                order.append((site, direc))
        lat, lon = ws.cell(row, 10).value, ws.cell(row, 11).value
        if isinstance(lat, (int, float)) and isinstance(lon, (int, float)):
            gps[site] = (float(lat), float(lon))
        if not _installed(ws.cell(row, 7).value):
            continue
        if direc not in {"n", "e", "s", "w"}:
            continue
        tag = f"{site}{direc}"
        for key in serial_keys(ws.cell(row, 3).value):
            serial_to_tag[key] = tag
    return {"installed": installed, "serial_to_tag": serial_to_tag, "order": order, "gps": gps}


def match_day1(places: list[dict], sites: dict[str, dict], hose: dict[str, dict]) -> list[dict]:
    by_street: dict[str, list[str]] = {}
    for sid, info in sites.items():
        by_street.setdefault(street_key(info["street"]), []).append(sid)
    rows = []
    used: set[str] = set()
    for pl in places:
        key = street_key(pl["name"])
        pool = by_street.get(key)
        if pool is None:
            pool = []
            for cand, sids in by_street.items():
                if abs(len(cand) - len(key)) <= 1 and (cand.startswith(key[:6]) or key.startswith(cand[:6])):
                    pool = pool + sids
        cands = [s for s in pool if s not in used]
        if len(cands) != 1:
            raise ValueError(f"{pl['name']!r} street {key!r} matched {cands}")
        sid = cands[0]
        seg = sites[sid]
        dist, along = point_to_segment_m(
            pl["lat"], pl["lon"],
            float(seg["begin_lat"]), float(seg["begin_lon"]),
            float(seg["end_lat"]), float(seg["end_lon"]),
        )
        if dist > 400:
            raise ValueError(f"{sid} {pl['name']} is {dist:.0f} m off the segment")
        used.add(sid)
        file = hose.get(sid)
        direc = file["dir"] if file else ""
        if not direc:
            inferred = infer_from_segment(
                float(seg["begin_lat"]), float(seg["begin_lon"]),
                float(seg["end_lat"]), float(seg["end_lon"]),
            )
            direc = inferred.get("direction") or ""
        rows.append({
            "site": sid,
            "street": seg["street"],
            "name": pl["name"],
            "lat": pl["lat"],
            "lon": pl["lon"],
            "dist_m": round(dist, 1),
            "on_segment": 0.02 <= along <= 0.98,
            "serial": file["serial4"] if file else "",
            "dir": direc,
            "installed": file is not None,
        })
    missing = sorted(set(hose) - used)
    if missing:
        raise ValueError(f"hose files with no Apple street: {missing}")
    return rows


def load_hoses(folder: Path) -> tuple[dict[str, dict], list[str]]:
    """Site -> hose info. A second file with the same study is a duplicate."""
    found: dict[str, dict] = {}
    notes: list[str] = []
    files = []
    for path in sorted(folder.glob("*.tvp")):
        serial, loc = tvp_header(path)
        m = re.match(r"(\d+)([nesw])c1b", loc, re.I)
        if not m:
            raise ValueError(f"{path.name}: bad location {loc}")
        site, direc = m.group(1), m.group(2).lower()
        start = __import__("core.picocount_hits", fromlist=["decode_start_time_tvp"]).decode_start_time_tvp(
            path.read_bytes()
        )
        files.append({
            "path": path,
            "site": site,
            "dir": direc,
            "serial": serial,
            "serial4": serial[-4:],
            "start": start,
            "header": loc,
        })
    by_study: dict[tuple, list[dict]] = {}
    for item in files:
        by_study.setdefault((item["serial"], item["start"], item["header"]), []).append(item)
    for study, group in by_study.items():
        keep = sorted(group, key=lambda g: (g["path"].name != g["header"] + ".tvp", g["path"].name))[0]
        for extra in group:
            if extra["path"] != keep["path"]:
                notes.append(
                    f"duplicate {extra['path'].name} is the same study as {keep['path'].name} "
                    f"(serial {study[0]}, start {study[1]})"
                )
        if keep["site"] in found:
            raise ValueError(f"two hoses for site {keep['site']}")
        found[keep["site"]] = keep
    return found, notes


def write_day1_sheet(tfc: Path, rows: list[dict]) -> None:
    wb = openpyxl.load_workbook(tfc)
    if "Day 1" in wb.sheetnames:
        del wb["Day 1"]
    src = wb["Day 2"]
    ws = wb.create_sheet("Day 1", 0)
    for col in range(1, 12):
        cell = ws.cell(1, col, src.cell(1, col).value)
        cell.font = copy(src.cell(1, col).font)
        cell.alignment = copy(src.cell(1, col).alignment)
        letter = openpyxl.utils.get_column_letter(col)
        ws.column_dimensions[letter].width = src.column_dimensions[letter].width
    stamped = datetime(DAY1_START.year, DAY1_START.month, DAY1_START.day)
    for i, row in enumerate(rows, start=2):
        ws.cell(i, 1, stamped).number_format = "M/D/YYYY"
        ws.cell(i, 2, row["site"])
        ws.cell(i, 3, row["serial"] or None)
        ws.cell(i, 4, row["dir"] or None)
        ws.cell(i, 5, 2)
        ws.cell(i, 6, None)
        ws.cell(i, 7, "x" if row["installed"] else None)
        ws.cell(i, 8, None)
        ws.cell(i, 9, "x" if row["installed"] else None)
        ws.cell(i, 10, row["lat"])
        ws.cell(i, 11, row["lon"])
    # Day 2 study date. Leave blank date cells blank.
    day2 = datetime(DAY2_START.year, DAY2_START.month, DAY2_START.day)
    for row in range(2, src.max_row + 1):
        if not _cell_id(src.cell(row, 2).value):
            continue
        if src.cell(row, 1).value in (None, ""):
            continue
        cell = src.cell(row, 1, day2)
        cell.number_format = "M/D/YYYY"
    wb.save(tfc)


def site_from_name(name: str) -> str:
    m = re.match(r"(\d+)", name)
    return m.group(1) if m else ""


def main() -> int:
    folder = FOLDER
    tfc = folder / TFC_NAME
    lines: list[str] = []

    def log(msg: str = "") -> None:
        print(msg)
        lines.append(msg)

    if not folder.is_dir() or not tfc.is_file() or not XLS.is_file():
        print("missing week 29 folder, TFC, or Isaacx", file=sys.stderr)
        return 1
    if not GUIDE_JS.is_file():
        print(f"missing Apple guide cache {GUIDE_JS}", file=sys.stderr)
        return 1

    places = apple_places()
    sites = ingest.parse_excel_sites([str(XLS)])
    hoses, dup_notes = load_hoses(folder)
    day2 = load_day2(tfc)
    # 4112, 4234, and 4298 are Day 2 hoses. They are not on the Day 1 guide.
    day1_hoses = {k: v for k, v in hoses.items() if k not in day2["installed"]}
    day1_rows = match_day1(places, sites, day1_hoses)

    backup = folder / BACKUP_NAME
    backup.mkdir(exist_ok=True)
    errors: list[str] = []

    log("Week 29")
    log(f"Day 1: {DAY1_START} .. {DAY1_END}  Day 2: {DAY2_START} .. {DAY2_END}")
    log("Keep each file's clock. Crop at the end date and that same clock.")
    log(f"Apple pins: {len(places)}  Day 1 rows: {len(day1_rows)}")
    for note in dup_notes:
        log(f"  {note}")
    log()
    for row in day1_rows:
        flag = "installed" if row["installed"] else "no count file"
        off = "" if row["on_segment"] else " off-segment"
        log(
            f"  {row['site']} {row['dir'] or '-'} serial {row['serial'] or '-'} "
            f"{row['dist_m']} m{off}  {row['name']}  {flag}"
        )
    log()

    w26.backup_file(tfc, backup)
    for src in list(folder.glob("*.tvp")) + list(folder.glob("*.csv")):
        if src.is_file() and src.parent == folder:
            w26.backup_file(src, backup)
    for note_item in dup_notes:
        # "duplicate NAME is the same study"
        name = note_item.split()[1]
        src = folder / name
        if src.exists():
            src.unlink()
            log(f"removed duplicate from the work folder: {name}")
    log()

    try:
        write_day1_sheet(tfc, day1_rows)
        log(f"TFC Day 1 sheet: {len(day1_rows)} rows, date {DAY1_START}")
        log(f"TFC Day 2 date cells -> {DAY2_START}")
    except Exception as exc:  # noqa: BLE001
        errors.append(f"TFC: {exc}")
        log(f"FAIL TFC: {exc}")
        bak = backup / tfc.name
        if bak.exists():
            shutil.copy2(bak, tfc)
        report = folder / "WEEK29_REPORT.txt"
        report.write_text("\n".join(lines) + "\n", encoding="utf-8")
        return 1

    day1_ids = {r["site"] for r in day1_rows if r["installed"]}
    day2_ids = day2["installed"]
    n_ok = 0

    log()
    log("=== TVP ===")
    for src in sorted(folder.glob("*.tvp")):
        try:
            _serial, loc = tvp_header(src)
        except Exception as exc:  # noqa: BLE001
            errors.append(f"{src.name}: {exc}")
            log(f"  FAIL {src.name}: {exc}")
            continue
        site = re.match(r"(\d+)", loc).group(1)
        if site in day1_ids:
            start_day, end_day = DAY1_START, DAY1_END
        elif site in day2_ids:
            start_day, end_day = DAY2_START, DAY2_END
        else:
            errors.append(f"TVP site not on TFC: {src.name}")
            log(f"  FAIL TVP site not on TFC: {src.name}")
            continue
        try:
            w26.process_tvp(src, src, start_day=start_day, end_day=end_day, lines=lines)
            n_ok += 1
        except Exception as exc:  # noqa: BLE001
            errors.append(f"{src.name}: {exc}")
            log(f"  FAIL {src.name}: {exc}")
            bak = backup / src.name
            if bak.exists():
                shutil.copy2(bak, src)

    log()
    log("=== volume ===")
    for src in sorted(folder.glob("*.csv")):
        if ".60min" in src.name.lower():
            continue
        site = site_from_name(src.name)
        if site not in day2_ids:
            errors.append(f"volume site not on Day 2: {src.name}")
            log(f"  FAIL volume site not on Day 2: {src.name}")
            continue
        try:
            w26.process_volume(src, src, lines)
            n_ok += 1
        except Exception as exc:  # noqa: BLE001
            errors.append(f"{src.name}: {exc}")
            log(f"  FAIL {src.name}: {exc}")
            bak = backup / src.name
            if bak.exists():
                shutil.copy2(bak, src)

    log()
    log("=== 60min ===")
    for src in sorted(folder.glob("*.60min.csv")):
        parsed = parse_60min(src)
        if not parsed["rows"]:
            log(f"  SKIP empty {src.name}")
            continue
        try:
            w26.process_60min(src, src, lines)
            n_ok += 1
        except Exception as exc:  # noqa: BLE001
            errors.append(f"{src.name}: {exc}")
            log(f"  FAIL {src.name}: {exc}")
            bak = backup / src.name
            if bak.exists():
                shutil.copy2(bak, src)

    log()
    log("=== 60min names ===")
    for src in sorted(folder.glob("*.60min.csv")):
        parsed = parse_60min(src)
        if not parsed["rows"]:
            continue
        new_name = dest_name(src.name, day2["serial_to_tag"])
        if new_name is None:
            errors.append(f"no TFC match: {src.name}")
            log(f"  FAIL no TFC match: {src.name}")
            continue
        if new_name == src.name:
            log(f"  already {src.name}")
            continue
        dest = src.with_name(new_name)
        if dest.exists():
            errors.append(f"dest exists: {src.name} -> {new_name}")
            log(f"  FAIL dest exists: {new_name}")
            continue
        src.rename(dest)
        log(f"  {src.name} -> {new_name}")

    log()
    log("=== photos ===")
    photos = list_photos(folder)
    photo_backup = folder / PHOTO_BACKUP
    photo_backup.mkdir(exist_ok=True)
    assigned, photo_errors = _assign(photos, day2["order"])
    if photo_errors:
        errors.extend(photo_errors)
        for e in photo_errors:
            log(f"  FAIL {e}")
    else:
        staged = []
        for src, sid, dt in assigned:
            dest_bak = photo_backup / src.name
            if not dest_bak.exists():
                shutil.copy2(src, dest_bak)
            im = Image.open(src)
            im = ImageOps.exif_transpose(im)
            if im.mode != "RGB":
                im = im.convert("RGB")
            exif = im.getexif()
            exif[0x0112] = 1
            taken = datetime.combine(DAY2_START, dt.time())
            stamp = taken.strftime("%Y:%m:%d %H:%M:%S")
            exif[0x0132] = stamp
            sub = exif.get_ifd(0x8769)
            sub[0x9003] = stamp
            sub[0x9004] = stamp
            data = compress_about_800kb(im, exif.tobytes())
            staged.append((src, sid, dt, data))
        temps = []
        for src, sid, dt, data in staged:
            tmp = src.with_name(f"_w29_{src.stem}.jpg")
            src.rename(tmp)
            temps.append((tmp, sid, dt, data))
        used: dict[str, int] = {}
        for src, sid, dt, data in temps:
            src.write_bytes(data)
            taken = datetime.combine(DAY2_START, dt.time())
            set_file_times(src, taken)
            check = read_datetime_original(src)
            if check.date() != DAY2_START or check.time() != dt.time():
                errors.append(f"{src.name}: Date Taken {check} want {taken}")
                log(f"  FAIL {src.name}: {check}")
                continue
            size = src.stat().st_size
            if not (LOW <= size <= HIGH):
                errors.append(f"{sid} photo {size} bytes")
            used[sid] = used.get(sid, 0) + 1
            suffix = "" if used[sid] == 1 else f"-{used[sid]}"
            new_name = f"{sid}{suffix}.jpg"
            dest = src.with_name(new_name)
            if dest.exists():
                errors.append(f"photo dest exists: {new_name}")
                log(f"  FAIL photo dest exists: {new_name}")
                continue
            src.rename(dest)
            log(f"  -> {new_name}: {size / 1024:.0f} KB  taken {check}")

    log()
    log("=== EST ===")
    gps1 = {r["site"]: (r["lat"], r["lon"]) for r in day1_rows}
    if EST_DAY1.is_file():
        est_backup = EST_DAY1.parent / "_week29_est_backup"
        est_backup.mkdir(exist_ok=True)
        pin_ids = {p["site"] for p in pushpins_from_est(EST_DAY1)}
        subset = {s: gps1[s] for s in gps1 if s in pin_ids}
        for err in patch_map(EST_DAY1, est_backup, subset, lines):
            errors.append(err)
            log(f"  FAIL {err}")
    map2 = folder / "Week 29 Map 2.est"
    if map2.is_file():
        pin_ids = {p["site"] for p in pushpins_from_est(map2)}
        subset = {s: day2["gps"][s] for s in day2["gps"] if s in pin_ids}
        for err in patch_map(map2, backup, subset, lines):
            errors.append(err)
            log(f"  FAIL {err}")

    log()
    log(f"processed ok: {n_ok}")
    log(f"errors: {len(errors)}")
    for e in errors:
        log(f"  ERROR {e}")
    report = folder / "WEEK29_REPORT.txt"
    report.write_text("\n".join(lines) + "\n", encoding="utf-8")
    log(f"Report: {report}")
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
