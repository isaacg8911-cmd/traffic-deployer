"""Week 27: write TFC field GPS onto Map 1.est (Day 2) and Map 2.est (Day 1).

Does not restamp counts. Backs up EST files to _week27_backup/ on first run.
"""
from __future__ import annotations

import math
import shutil
import sys
from pathlib import Path

import openpyxl

APP = Path(__file__).resolve().parents[1]
SCRIPTS = Path(__file__).resolve().parent
sys.path.insert(0, str(APP))
sys.path.insert(0, str(SCRIPTS))

from core.est_field_gps import apply_field_gps_to_est, find_paired_inline_coords  # noqa: E402
from core.est_viewer import _SITE_ID, pushpins_from_est  # noqa: E402
from week27_demand import (  # noqa: E402
    DAY1_SHEET,
    DAY2_SHEET,
    FOLDER,
    TFC_NAME,
    _cell_id,
    _sheet,
)

MAP1 = "Week 27 Map 1.est"  # Day 2 10xxx
MAP2 = "Week 27 Map 2.est"  # Day 1 4xxx
BACKUP_NAME = "_week27_backup"
MAX_VERIFY_M = 25.0


def hav_m(lat1, lon1, lat2, lon2) -> float:
    r = 6371000.0
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp = math.radians(lat2 - lat1)
    dl = math.radians(lon2 - lon1)
    a = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * r * math.asin(math.sqrt(a))


def load_tfc_gps(tfc: Path, sheet: str, day: int) -> dict[str, tuple[float, float]]:
    wb = openpyxl.load_workbook(tfc, data_only=True)
    ws = _sheet(wb, sheet, day=day)
    out: dict[str, tuple[float, float]] = {}
    for row in range(2, ws.max_row + 1):
        site = _cell_id(ws.cell(row, 2).value)
        lat, lon = ws.cell(row, 10).value, ws.cell(row, 11).value
        if not site or not isinstance(lat, (int, float)) or not isinstance(lon, (int, float)):
            continue
        out[site] = (float(lat), float(lon))
    return out


def backup_est(src: Path, backup: Path) -> Path:
    dest = backup / src.name
    if not dest.exists():
        shutil.copy2(src, dest)
    return dest


def all_inline_pins(est: Path) -> dict[str, list[dict]]:
    """Every pushpin, including begin/end duplicates (viewer keeps first only)."""
    data = est.read_bytes()
    out: dict[str, list[dict]] = {}
    pos = 0
    while True:
        pos = data.find(b"\xff\xfe", pos)
        if pos < 0:
            break
        j = pos + 2
        digits: list[str] = []
        while j < len(data) and 48 <= data[j] <= 57:
            digits.append(chr(data[j]))
            j += 1
        site = "".join(digits)
        if not _SITE_ID.match(site):
            pos += 2
            continue
        tail = data[pos : pos + 400]
        pair = find_paired_inline_coords(tail)
        if not pair:
            pos += 2
            continue
        la0, la1, lo0, lo1 = pair
        rec = {"site": site, "lat": float(tail[la0:la1]), "lon": float(tail[lo0:lo1])}
        out.setdefault(site, []).append(rec)
        pos += 2
    return out


def verify_pins(
    est: Path, want: dict[str, tuple[float, float]]
) -> tuple[list[str], list[tuple[str, float]], list[tuple[str, float]]]:
    grouped = all_inline_pins(est)
    if not grouped:
        grouped = {p["site"]: [p] for p in pushpins_from_est(est)}
    missing: list[str] = []
    far: list[tuple[str, float]] = []
    ok: list[tuple[str, float]] = []
    for sid, (lat, lon) in sorted(want.items(), key=lambda kv: (len(kv[0]), kv[0])):
        pins = grouped.get(sid) or []
        if not pins:
            missing.append(sid)
            continue
        d = min(hav_m(lat, lon, p["lat"], p["lon"]) for p in pins)
        if d > MAX_VERIFY_M:
            far.append((sid, d))
        else:
            ok.append((sid, d))
    return missing, far, ok


def patch_map(est: Path, backup: Path, updates: dict[str, tuple[float, float]], lines: list[str]) -> list[str]:
    errors: list[str] = []
    if not updates:
        errors.append(f"{est.name}: no TFC GPS to write")
        return errors
    before = {p["site"]: p for p in pushpins_from_est(est)}
    backup_est(est, backup)
    result = apply_field_gps_to_est(est, updates, est)
    lines.append(
        f"{est.name}: patched {result['sites_patched']}/{result['sites_requested']} "
        f"pins={result['pins_patched']} fmt={result['format']}"
    )
    if result["missing_in_est"]:
        errors.append(f"{est.name} missing pins: {result['missing_in_est']}")
    if result["warnings"]:
        for w in result["warnings"]:
            lines.append(f"  warn {w}")
    missing, far, ok = verify_pins(est, updates)
    moved = 0
    after = {p["site"]: p for p in pushpins_from_est(est)}
    for sid in updates:
        b, a = before.get(sid), after.get(sid)
        if b and a and (abs(b["lat"] - a["lat"]) > 1e-6 or abs(b["lon"] - a["lon"]) > 1e-6):
            moved += 1
    lines.append(f"  verify ok={len(ok)} moved={moved} far={len(far)} missing={len(missing)}")
    for sid, d in far[:12]:
        errors.append(f"{est.name} site {sid} {d:.1f}m from TFC")
    for sid in missing:
        errors.append(f"{est.name} site {sid} not in EST after patch")
    if result["sites_patched"] < 1:
        errors.append(f"{est.name}: patched 0 sites")
    return errors


def main() -> int:
    folder = Path(sys.argv[1]) if len(sys.argv) > 1 else FOLDER
    tfc = folder / TFC_NAME
    if not folder.is_dir():
        print(f"missing folder {folder}", file=sys.stderr)
        return 1
    if not tfc.is_file():
        print(f"missing TFC {tfc}", file=sys.stderr)
        return 1

    map1 = folder / MAP1
    map2 = folder / MAP2
    for p in (map1, map2):
        if not p.is_file():
            print(f"missing EST {p}", file=sys.stderr)
            return 1

    backup = folder / BACKUP_NAME
    backup.mkdir(exist_ok=True)
    day2 = load_tfc_gps(tfc, DAY2_SHEET, 2)
    day1 = load_tfc_gps(tfc, DAY1_SHEET, 1)
    lines: list[str] = []
    errors: list[str] = []

    print(f"Folder: {folder}")
    print(f"TFC Day 2 GPS: {len(day2)}  Day 1 GPS: {len(day1)}")
    errors.extend(patch_map(map1, backup, day2, lines))
    errors.extend(patch_map(map2, backup, day1, lines))
    for msg in lines:
        print(msg)

    report = folder / "WEEK27_EST_GPS.txt"
    report.write_text("\n".join(lines + [""] + [f"ERROR {e}" for e in errors]) + "\n", encoding="utf-8")
    print(f"Report: {report}")
    if errors:
        for e in errors:
            print(f"  ERROR {e}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
