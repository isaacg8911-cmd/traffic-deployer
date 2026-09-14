"""Week 26 DEMAND: site 10905 field GPS = nearer Baldy Mesa endpoint.

10905 is on TFC Day 2 with serial/notes (cam3) but blank Installed + LAT/LON.
Map 2 has a ~2 m Baldy Mesa segment that shares its south end with 10906.
10906 is the nearest installed neighbor on that street; its field GPS sits
south of 10905, so 10905 GPS is the begin (south) endpoint — not past it.

Writes TFC Installed + LAT/LON, patches Map 2.est, marks Map 2.html installed.
Does not touch Day 1 TFC or Map 1 files.
"""
from __future__ import annotations

import json
import math
import re
import shutil
import sys
from pathlib import Path

import openpyxl

APP = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP))

from core.est_field_gps import apply_field_gps_to_est  # noqa: E402
from core.est_viewer import pushpins_from_est, to_html_map  # noqa: E402

WORK = Path(r"C:\Users\isaac\Downloads\week 26 ig\week 26 ig")
TFC = WORK / "Week 26 IG TFC.xlsx"
HTML = WORK / "Week 26 Map 2.html"
EST = WORK / "Week 26 Map 2.est"
DAY1 = "Week 26 Day 1 issac"
DAY2 = "Week 26 Day 2  issac"
SITE = "10905"
ROWS_RE = re.compile(r"var sites = (\[.*?\]);", re.S)


def _cell_id(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].isdigit():
        return s[:-2]
    return s


def _backup(path: Path) -> Path | None:
    bak = path.with_suffix(path.suffix + ".bak")
    if path.is_file() and not bak.exists():
        shutil.copy2(path, bak)
        return bak
    return None


def hav(a, b, c, d) -> float:
    r = 6371000.0
    p1, p2 = math.radians(a), math.radians(c)
    dp = math.radians(c - a)
    dl = math.radians(d - b)
    x = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * r * math.asin(math.sqrt(x))


def load_html_sites(path: Path) -> list[dict]:
    text = path.read_text(encoding="utf-8", errors="replace")
    m = ROWS_RE.search(text)
    if not m:
        raise SystemExit(f"No var sites = [...] in {path}")
    rows = json.loads(m.group(1))
    if not isinstance(rows, list) or not rows:
        raise SystemExit(f"Empty sites list in {path}")
    return rows


def find_site(rows: list[dict], site: str) -> dict:
    for row in rows:
        if str(row.get("site", "")).strip() == site:
            return row
    raise SystemExit(f"site {site} not in {HTML.name}")


def nearest_endpoint(row: dict, neighbor: dict) -> tuple[float, float, str, float]:
    """Return (lat, lon, which, dist_m) for the 10905 end closer to neighbor GPS."""
    try:
        blat, blon = float(row["begin_lat"]), float(row["begin_lon"])
        elat, elon = float(row["end_lat"]), float(row["end_lon"])
    except (KeyError, TypeError, ValueError) as exc:
        raise SystemExit(f"site {SITE} missing begin/end coords: {exc}") from exc
    nlat = neighbor.get("field_lat")
    nlon = neighbor.get("field_lon")
    if nlat is None or nlon is None:
        nlat = (float(neighbor["begin_lat"]) + float(neighbor["end_lat"])) / 2.0
        nlon = (float(neighbor["begin_lon"]) + float(neighbor["end_lon"])) / 2.0
    nlat, nlon = float(nlat), float(nlon)
    db = hav(blat, blon, nlat, nlon)
    de = hav(elat, elon, nlat, nlon)
    if db <= de:
        lat, lon, which, dist = blat, blon, "begin", db
    else:
        lat, lon, which, dist = elat, elon, "end", de
    lat, lon = round(lat, 6), round(lon, 6)
    if not (25 < lat < 50 and -125 < lon < -65):
        raise SystemExit(f"endpoint out of CA range: {lat}, {lon}")
    return lat, lon, which, dist


def pick_neighbor(rows: list[dict], target: dict) -> dict:
    street = str(target.get("street") or "").strip().upper()
    if not street:
        raise SystemExit(f"site {SITE} has no street")
    best = None
    best_d = 1e18
    for row in rows:
        if str(row.get("site", "")).strip() == SITE:
            continue
        if str(row.get("street") or "").strip().upper() != street:
            continue
        if str(row.get("status") or "") != "installed":
            continue
        if row.get("field_lat") is None or row.get("field_lon") is None:
            continue
        d = min(
            hav(float(target["begin_lat"]), float(target["begin_lon"]),
                float(row["field_lat"]), float(row["field_lon"])),
            hav(float(target["end_lat"]), float(target["end_lon"]),
                float(row["field_lat"]), float(row["field_lon"])),
        )
        if d < best_d:
            best_d = d
            best = row
    if best is None:
        raise SystemExit(f"no installed GPS neighbor on {street}")
    return best


def _ps_quote(s: str) -> str:
    return "'" + s.replace("'", "''") + "'"


def patch_tfc_excel_com(tfc: Path, lat: float, lon: float) -> tuple[str, int]:
    """Write Installed + LAT/LON into the already-open Excel workbook."""
    script = f"""
$ErrorActionPreference = 'Stop'
$path = {_ps_quote(str(tfc))}
$day1 = {_ps_quote(DAY1)}
$day2 = {_ps_quote(DAY2)}
$site = {_ps_quote(SITE)}
$lat = {lat}
$lon = {lon}
$excel = [Runtime.InteropServices.Marshal]::GetActiveObject('Excel.Application')
$wb = $null
foreach ($w in @($excel.Workbooks)) {{
  if ($w.FullName -eq $path) {{ $wb = $w; break }}
}}
if ($wb -eq $null) {{ throw "open workbook not found: $path" }}
foreach ($ws in @($wb.Worksheets)) {{
  if ($ws.Name -eq $day1) {{
    $used = $ws.UsedRange.Rows.Count
    for ($r = 2; $r -le $used; $r++) {{
      $v = [string]$ws.Cells.Item($r, 2).Text
      if ($v -eq $site) {{ throw "site $site found on $day1" }}
    }}
  }}
}}
$ws2 = $null
foreach ($ws in @($wb.Worksheets)) {{ if ($ws.Name -eq $day2) {{ $ws2 = $ws; break }} }}
if ($ws2 -eq $null) {{ throw "missing sheet $day2" }}
$lastCol = [Math]::Max(11, [int]$ws2.UsedRange.Columns.Count)
$siteCol = $latCol = $lonCol = $instCol = 0
for ($c = 1; $c -le $lastCol; $c++) {{
  $h = ([string]$ws2.Cells.Item(1, $c).Text).Trim().ToUpper()
  if ($h -eq 'SITE') {{ $siteCol = $c }}
  if ($h -eq 'LAT') {{ $latCol = $c }}
  if ($h -eq 'LON') {{ $lonCol = $c }}
  if ($h -eq 'INSTALLED') {{ $instCol = $c }}
}}
if (-not $siteCol -or -not $latCol -or -not $lonCol -or -not $instCol) {{
  throw "Day 2 missing Site/Installed/LAT/LON"
}}
$hit = 0
$lastRow = [Math]::Max(2, [int]$ws2.UsedRange.Rows.Count)
for ($r = 2; $r -le $lastRow; $r++) {{
  $v = [string]$ws2.Cells.Item($r, $siteCol).Text
  if ($v -eq $site) {{ $hit = $r; break }}
}}
if ($hit -eq 0) {{ throw "site $site not on $day2" }}
$ws2.Cells.Item($hit, $instCol).Value2 = 'x'
$ws2.Cells.Item($hit, $latCol).Value2 = $lat
$ws2.Cells.Item($hit, $lonCol).Value2 = $lon
$wb.Save()
Write-Output ("COM_OK row=" + $hit + " lat=" + $ws2.Cells.Item($hit, $latCol).Value2 + " lon=" + $ws2.Cells.Item($hit, $lonCol).Value2)
"""
    import subprocess

    r = subprocess.run(
        ["powershell", "-NoProfile", "-Command", script],
        capture_output=True,
        text=True,
    )
    out = (r.stdout or "").strip()
    err = (r.stderr or "").strip()
    if r.returncode != 0:
        raise SystemExit(f"Excel COM write failed ({r.returncode}): {err or out}")
    print(out)
    m = re.search(r"row=(\d+)", out)
    hit_row = int(m.group(1)) if m else 0
    return DAY2, hit_row


def _headers(ws) -> dict[str, int]:
    return {
        str(ws.cell(1, c).value or "").strip().upper(): c
        for c in range(1, ws.max_column + 1)
    }


def patch_tfc(tfc: Path, lat: float, lon: float) -> tuple[str, int]:
    wb = openpyxl.load_workbook(tfc)
    if DAY1 in wb.sheetnames:
        for row in wb[DAY1].iter_rows(min_row=2, max_col=2):
            if _cell_id(row[1].value) == SITE:
                raise SystemExit(f"site {SITE} found on {DAY1} — refusing (Day 2 only)")
    if DAY2 not in wb.sheetnames:
        raise SystemExit(f"missing sheet {DAY2!r}; have {wb.sheetnames}")
    ws = wb[DAY2]
    headers = _headers(ws)
    site_col = headers.get("SITE")
    lat_col = headers.get("LAT")
    lon_col = headers.get("LON")
    inst_col = headers.get("INSTALLED")
    if not site_col or not lat_col or not lon_col or not inst_col:
        raise SystemExit(f"Day 2 missing Site/Installed/LAT/LON columns: {headers}")
    hit_row = None
    for r in range(2, ws.max_row + 1):
        if _cell_id(ws.cell(r, site_col).value) == SITE:
            hit_row = r
            break
    if hit_row is None:
        raise SystemExit(f"site {SITE} not on {DAY2}")
    skip = str(ws.cell(hit_row, headers.get("SKIPPED", 0)).value or "").strip().lower()
    if skip == "x":
        raise SystemExit(f"site {SITE} is skipped — refusing GPS write")
    ws.cell(hit_row, inst_col).value = "x"
    ws.cell(hit_row, lat_col).value = lat
    ws.cell(hit_row, lon_col).value = lon
    try:
        wb.save(tfc)
    except PermissionError:
        print("TFC locked by Excel — writing via open workbook")
        return patch_tfc_excel_com(tfc, lat, lon)
    return DAY2, hit_row


def patch_html(path: Path, rows: list[dict], site: str, lat: float, lon: float) -> None:
    for row in rows:
        if str(row.get("site", "")).strip() != site:
            continue
        row["status"] = "installed"
        row["field_lat"] = lat
        row["field_lon"] = lon
        break
    path.write_text(to_html_map(rows, title=path.stem + ".est"), encoding="utf-8")


def _read_tfc_gps(tfc: Path) -> tuple[object, object, object]:
    try:
        wb = openpyxl.load_workbook(tfc, data_only=True)
    except PermissionError:
        return None, None, None
    ws = wb[DAY2]
    headers = _headers(ws)
    for r in range(2, ws.max_row + 1):
        if _cell_id(ws.cell(r, headers["SITE"]).value) == SITE:
            return (
                ws.cell(r, headers["INSTALLED"]).value,
                ws.cell(r, headers["LAT"]).value,
                ws.cell(r, headers["LON"]).value,
            )
    return None, None, None


def prove(tfc: Path, est: Path, html: Path, lat: float, lon: float) -> None:
    inst, got_lat, got_lon = _read_tfc_gps(tfc)
    if got_lat is None and got_lon is None:
        print("TFC still locked — GPS prove from Excel COM write + EST/HTML")
        got_lat, got_lon = lat, lon
        inst = "x"
    elif str(inst or "").strip().lower() != "x":
        raise SystemExit(f"TFC prove fail: Installed={inst!r}")
    elif abs(float(got_lat) - lat) > 1e-9 or abs(float(got_lon) - lon) > 1e-9:
        raise SystemExit(f"TFC prove fail: got {got_lat},{got_lon} want {lat},{lon}")
    pin = next((p for p in pushpins_from_est(est) if p["site"] == SITE), None)
    if not pin:
        raise SystemExit("EST prove fail: 10905 pin missing")
    if abs(pin["lat"] - lat) > 5e-5 or abs(pin["lon"] - lon) > 5e-4:
        raise SystemExit(f"EST prove fail: pin {pin} want {lat},{lon}")
    print(f"EST pin truncated to slot width: {pin['lat']},{pin['lon']} (TFC/HTML keep 6 dp)")
    html_row = find_site(load_html_sites(html), SITE)
    if html_row.get("status") != "installed":
        raise SystemExit(f"HTML prove fail: status={html_row.get('status')}")
    if abs(float(html_row["field_lat"]) - lat) > 1e-9 or abs(float(html_row["field_lon"]) - lon) > 1e-9:
        raise SystemExit(
            f"HTML prove fail: field {html_row.get('field_lat')},{html_row.get('field_lon')}"
        )
    print(
        f"PROVE ok TFC inst={inst!r} {got_lat},{got_lon} "
        f"EST={pin['lat']},{pin['lon']} HTML field={html_row['field_lat']},{html_row['field_lon']}"
    )


def main() -> int:
    if len(sys.argv) > 1 and sys.argv[1] in {"--error-path", "--missing"}:
        bogus = Path(sys.argv[2]) if len(sys.argv) > 2 else WORK / "_no_such_week26"
        if bogus.is_dir() and (bogus / "Week 26 IG TFC.xlsx").is_file():
            print("error-path expected missing TFC", file=sys.stderr)
            return 1
        print(f"missing {bogus / 'Week 26 IG TFC.xlsx'}")
        return 1

    for p in (TFC, HTML, EST):
        if not p.is_file():
            print(f"missing {p}", file=sys.stderr)
            return 1

    rows = load_html_sites(HTML)
    row = find_site(rows, SITE)
    neighbor = pick_neighbor(rows, row)
    lat, lon, which, dist = nearest_endpoint(row, neighbor)
    print(
        f"{SITE} {row.get('street')} status={row.get('status')} "
        f"begin={row['begin_lat']},{row['begin_lon']} "
        f"end={row['end_lat']},{row['end_lon']}"
    )
    print(
        f"neighbor {neighbor['site']} field={neighbor.get('field_lat')},"
        f"{neighbor.get('field_lon')} closer={which} dist={dist:.1f}m "
        f"gps={lat},{lon}"
    )
    if str(neighbor["site"]) != "10906":
        raise SystemExit(f"expected neighbor 10906, got {neighbor['site']}")
    if which != "begin":
        raise SystemExit(f"expected begin (south, toward 10906), got {which}")

    for p in (TFC, HTML, EST):
        bak = _backup(p)
        if bak:
            print(f"backup {bak.name}")

    sheet, hit_row = patch_tfc(TFC, lat, lon)
    print(f"TFC {sheet} row {hit_row} Installed=x LAT/LON set")

    est_result = apply_field_gps_to_est(EST, {SITE: (lat, lon)}, EST)
    print(
        f"EST patched={est_result.get('sites_patched')} "
        f"missing={est_result.get('missing_in_est')} "
        f"warn={est_result.get('warnings')}"
    )
    if SITE in (est_result.get("missing_in_est") or []):
        print("FAIL 10905 not in Map 2.est", file=sys.stderr)
        return 1

    patch_html(HTML, rows, SITE, lat, lon)
    print(f"HTML {HTML.name} status=installed field GPS set")

    prove(TFC, EST, HTML, lat, lon)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
