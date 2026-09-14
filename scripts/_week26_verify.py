"""Verify Week 26 IG restamp: dates, times, TFC names. Exit 1 on failure."""
from __future__ import annotations

import sys
from datetime import date, datetime, timedelta
from pathlib import Path

APP = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import openpyxl  # noqa: E402
from core import picocount_hits  # noqa: E402
from week24_trim import locate_stream, parse_60min, parse_volume_csv  # noqa: E402
from week25_day2_rename_day1_shift import file_serial_key  # noqa: E402
from week26_demand import (  # noqa: E402
    DAY1_END,
    DAY1_START,
    DAY2_END,
    DAY2_START,
    load_tfc_sites,
    site_of,
)

FOLDER = Path(r"C:\Users\isaac\Downloads\week 26 ig\week 26 ig")


def _as_date(v) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return datetime.fromisoformat(str(v)[:10]).date()


def main() -> int:
    folder = Path(sys.argv[1]) if len(sys.argv) > 1 else FOLDER
    tfc = folder / "Week 26 IG TFC.xlsx"
    backup = folder / "_week26_backup"
    errs: list[str] = []

    wb = openpyxl.load_workbook(tfc, data_only=True)
    print("sheets", wb.sheetnames)
    expect = {wb.sheetnames[0]: DAY1_START, wb.sheetnames[1]: DAY2_START}
    for name, want in expect.items():
        ws = wb[name]
        dates: set[date] = set()
        n = 0
        for row in range(2, ws.max_row + 1):
            v = ws.cell(row, 1).value
            if v in (None, ""):
                continue
            n += 1
            d = _as_date(v)
            dates.add(d)
            if d != want:
                errs.append(f"TFC {name} row {row} date {d} != {want}")
        print(f"TFC {name}: {n} dates {sorted(dates)}")

    day1, day2, serial_to_tag = load_tfc_sites(tfc)
    print(f"sites d1={len(day1)} d2={len(day2)} serials={len(serial_to_tag)}")

    print("\n=== 60min ===")
    csv60 = sorted(folder.glob("*.60min.csv"))
    print("count", len(csv60))
    if len(csv60) != 15:
        errs.append(f"expected 15 60min CSVs, got {len(csv60)}")
    for p in csv60:
        d = parse_60min(p)
        first = d["rows"][0][0]
        last = d["rows"][-1][0]
        print(
            f"{p.name}: start={d['start']} end={d['end']} "
            f"rows={len(d['rows'])} first={first} last={last}"
        )
        if d["start"].date() != DAY2_START:
            errs.append(f"{p.name} start {d['start']}")
        if d["end"].date() != DAY2_END:
            errs.append(f"{p.name} end {d['end']} not on {DAY2_END}")
        if any(ts.date() > DAY2_END for ts, _ in d["rows"]):
            errs.append(f"{p.name} data row past {DAY2_END}")
        if any(ts.date() < DAY2_START for ts, _ in d["rows"]):
            errs.append(f"{p.name} data row before {DAY2_START}")
        key = file_serial_key(p.name)
        tag = serial_to_tag.get(key) if key is not None else None
        if not tag:
            errs.append(f"{p.name} serial {key} not on TFC")
        elif not p.name.startswith(f"{tag} "):
            errs.append(f"{p.name} missing prefix {tag}")
        bak = backup / p.name.split(" ", 1)[-1]
        if bak.exists():
            b = parse_60min(bak)
            if b["start"].time() != d["start"].time():
                errs.append(
                    f"{p.name} start time {d['start'].time()} != backup {b['start'].time()}"
                )
            if b["end"].time() != d["end"].time():
                errs.append(
                    f"{p.name} end time {d['end'].time()} != backup {b['end'].time()}"
                )

    print("\n=== volume ===")
    vols = [
        p
        for p in sorted(folder.glob("*.csv"))
        if ".60min" not in p.name.lower() and site_of(p.name) in day2
    ]
    if len(vols) != 5:
        errs.append(f"expected 5 volume CSVs, got {len(vols)}")
    for p in vols:
        d = parse_volume_csv(p)
        days = [str(x["date"]) for x in d["days"]]
        print(
            f"{p.name}: started={d['started']} ended={d['ended']} days={days} "
            f"n={sum(len(x['intervals']) for x in d['days'])}"
        )
        if d["started"].date() != DAY2_START:
            errs.append(f"{p.name} started {d['started']}")
        if d["ended"].date() != DAY2_END:
            errs.append(f"{p.name} ended {d['ended']}")
        if any(x["date"] > DAY2_END or x["date"] < DAY2_START for x in d["days"]):
            errs.append(f"{p.name} day out of window {days}")
        bak = backup / p.name
        if bak.exists():
            b = parse_volume_csv(bak)
            if b["started"].time() != d["started"].time():
                errs.append(f"{p.name} start time changed")
            if b["ended"].time() != d["ended"].time():
                errs.append(f"{p.name} end time changed")
            past = [x["date"] for x in b["days"] if x["date"] > DAY2_END]
            if not past:
                errs.append(f"{p.name} backup has no days after {DAY2_END} (trim unproven)")
            if any(x["date"] in past for x in d["days"]):
                errs.append(f"{p.name} still has post-window days")

    print("\n=== TVP ===")
    tvps = sorted(folder.glob("*.tvp"))
    d1n = d2n = 0
    for p in tvps:
        data = p.read_bytes()
        start = picocount_hits.decode_start_time_tvp(data)
        _off, hits, _end = locate_stream(data)
        valid = [h for h in hits if 0 <= h["seconds"] <= 86400 * 8]
        last = start + timedelta(seconds=valid[-1]["seconds"]) if start and valid else None
        site = site_of(p.name)
        if site in day1:
            want_s, want_e, kind = DAY1_START, DAY1_END, "D1"
            d1n += 1
        elif site in day2:
            want_s, want_e, kind = DAY2_START, DAY2_END, "D2"
            d2n += 1
        else:
            errs.append(f"TVP not on TFC {p.name}")
            want_s = want_e = None
            kind = "?"
        print(f"{p.name} {kind}: start={start} last={last} hits={len(valid)} bytes={len(data)}")
        if want_s and (start is None or start.date() != want_s):
            errs.append(f"{p.name} start {start} != {want_s}")
        if want_e and (last is None or last.date() > want_e):
            errs.append(f"{p.name} last {last} after {want_e}")
        bak = backup / p.name
        if bak.exists() and start:
            bdata = bak.read_bytes()
            bs = picocount_hits.decode_start_time_tvp(bdata)
            if bs and bs.time() != start.time():
                errs.append(f"{p.name} start time {start.time()} != backup {bs.time()}")
            if want_e:
                _bo, bhits, _be = locate_stream(bdata)
                bvalid = [h for h in bhits if 0 <= h["seconds"] <= 86400 * 8]
                if bs and bvalid:
                    blast = bs + timedelta(seconds=bvalid[-1]["seconds"])
                    if blast.date() > want_e and last and last.date() > want_e:
                        errs.append(f"{p.name} backup past {want_e} but output still {last}")
    if d1n != 21:
        errs.append(f"expected 21 Day 1 TVPs, got {d1n}")
    if d2n != 2:
        errs.append(f"expected 2 Day 2 TVPs, got {d2n}")

    print()
    print(f"=== ERRORS {len(errs)} ===")
    for e in errs:
        print(" ", e)
    return 1 if errs else 0


if __name__ == "__main__":
    raise SystemExit(main())
