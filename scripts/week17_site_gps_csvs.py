"""Week 17: one GPS CSV with street + begin/end/install for all sites.

Sources:
  - Week 17 Install Pins.html  (street + begin/end/install)
  - Week 17 IG TFC.xlsx        (cross-check install LAT/LON)

Writes:
  <work>/Week17_Sites_GPS.csv
"""
from __future__ import annotations

import csv
import json
import re
import shutil
import sys
from pathlib import Path

import openpyxl

APP = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP))

WORK = Path(r"C:\Users\isaac\Downloads\week 17 ig\week 17 ig")
HTML = WORK / "Week 17 Install Pins.html"
TFC = WORK / "Week 17 IG TFC.xlsx"
OUT = WORK / "Week17_Sites_GPS.csv"
OLD_DIR = WORK / "site_gps_csv"

ROWS_RE = re.compile(r"const rows = (\[.*?\]);", re.S)


def _site_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def _fmt(v) -> str:
    if v is None:
        return ""
    return f"{float(v):.6f}"


def load_html_rows(path: Path) -> list[dict]:
    text = path.read_text(encoding="utf-8", errors="replace")
    m = ROWS_RE.search(text)
    if not m:
        raise SystemExit(f"No const rows = [...] in {path}")
    rows = json.loads(m.group(1))
    if not isinstance(rows, list) or not rows:
        raise SystemExit(f"Empty rows list in {path}")
    return rows


def load_tfc_install(path: Path) -> dict[str, tuple[float, float]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    out: dict[str, tuple[float, float]] = {}
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i == 1 or not row or row[1] is None:
            continue
        site = _site_str(row[1])
        lat, lon = row[9], row[10]
        if isinstance(lat, (int, float)) and isinstance(lon, (int, float)):
            out[site] = (float(lat), float(lon))
    return out


def main() -> int:
    if not HTML.is_file():
        print(f"FAIL missing HTML: {HTML}")
        return 1
    if not TFC.is_file():
        print(f"FAIL missing TFC: {TFC}")
        return 1

    rows = load_html_rows(HTML)
    rows = sorted(rows, key=lambda r: (len(_site_str(r.get("site"))), _site_str(r.get("site"))))
    tfc = load_tfc_install(TFC)

    mismatches: list[str] = []
    with OUT.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([
            "site", "street",
            "begin_lat", "begin_lon",
            "end_lat", "end_lon",
            "install_lat", "install_lon",
        ])
        for row in rows:
            site = _site_str(row.get("site"))
            street = str(row.get("street") or "").strip()
            ilat, ilon = row.get("install_lat"), row.get("install_lon")
            if site in tfc and ilat is not None and ilon is not None:
                tlat, tlon = tfc[site]
                if abs(tlat - float(ilat)) > 1e-6 or abs(tlon - float(ilon)) > 1e-6:
                    mismatches.append(site)
            w.writerow([
                site, street,
                _fmt(row.get("begin_lat")), _fmt(row.get("begin_lon")),
                _fmt(row.get("end_lat")), _fmt(row.get("end_lon")),
                _fmt(ilat), _fmt(ilon),
            ])

    if OLD_DIR.is_dir():
        shutil.rmtree(OLD_DIR)

    print(f"OK wrote {OUT} sites={len(rows)}")
    if mismatches:
        print(f"WARN HTML vs TFC mismatch sites: {', '.join(mismatches)}")
        return 2
    print("OK HTML install matches TFC LAT/LON")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
